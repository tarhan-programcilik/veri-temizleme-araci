import io
from datetime import datetime
import pandas as pd
import streamlit as st
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Sayfa Yapılandırması
st.set_page_config(
    page_title="Veri Temizleme ve Özetleme Aracı",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Özel CSS ile daha modern ve şık görünüm
st.markdown("""
<style>
    .main-header {
        font-size: 2.2rem;
        font-weight: 700;
        color: #1E3A8A;
        margin-bottom: 0.2rem;
    }
    .sub-header {
        font-size: 1rem;
        color: #64748B;
        margin-bottom: 1.5rem;
    }
    .metric-card {
        background-color: #F8FAFC;
        border: 1px solid #E2E8F0;
        border-radius: 8px;
        padding: 12px 16px;
        text-align: center;
    }
    .stAlert {
        border-radius: 8px;
    }
</style>
""", unsafe_allow_html=True)


def log_action(message: str):
    """Yapılan işlemi zaman damgasıyla geçmişe kaydeder."""
    timestamp = datetime.now().strftime("%H:%M:%S")
    st.session_state.history_log.append(f"[{timestamp}] {message}")


def init_session_state():
    """Oturum durumlarını başlatır."""
    if "df" not in st.session_state:
        st.session_state.df = None
    if "raw_df" not in st.session_state:
        st.session_state.raw_df = None
    if "history_log" not in st.session_state:
        st.session_state.history_log = []
    if "file_id" not in st.session_state:
        st.session_state.file_id = None


init_session_state()


def to_excel_bytes(df: pd.DataFrame) -> bytes:
    """DataFrame'i stillendirilmiş openpyxl Excel bytes formatına dönüştürür."""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Temiz_Veri')
        worksheet = writer.sheets['Temiz_Veri']

        # Başlık biçimlendirmesi (Koyu mavi arka plan, beyaz kalın yazı)
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        # Otomatik sütun genişliği ayarlama
        for col_idx, col in enumerate(df.columns, 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(
                len(str(col)),
                df[col].astype(str).str.len().max() if not df.empty else 0
            )
            worksheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    return output.getvalue()


# -------------------------------------------------------------
# KENAR ÇUBUĞU (SIDEBAR) - DOSYA YÜKLEME & AYARLAR
# -------------------------------------------------------------
with st.sidebar:
    st.header("📂 Dosya Yükleme")
    uploaded_file = st.file_uploader(
        "Excel veya CSV dosyası seçin",
        type=["csv", "xlsx", "xls"],
        help="Maksimum 200MB desteklenir."
    )

    if uploaded_file is not None:
        file_identifier = f"{uploaded_file.name}_{uploaded_file.size}"
        
        # Yeni bir dosya yüklendiğinde oturumu yenile
        if st.session_state.file_id != file_identifier:
            st.session_state.file_id = file_identifier
            st.session_state.history_log = []
            
            try:
                if uploaded_file.name.endswith('.csv'):
                    # CSV için kodlama ve ayraç seçenekleri
                    st.subheader("CSV Ayarları")
                    encoding_opt = st.selectbox("Karakter Kodlaması", ["utf-8", "latin5", "iso-8859-9", "cp1254", "utf-8-sig"])
                    sep_opt = st.selectbox("Ayraç (Delimiter)", [",", ";", "\\t", "|"])
                    
                    df_loaded = pd.read_csv(uploaded_file, sep=sep_opt, encoding=encoding_opt)
                else:
                    # Excel dosyası - sayfa seçimi
                    excel_file = pd.ExcelFile(uploaded_file)
                    sheet_names = excel_file.sheet_names
                    selected_sheet = st.selectbox("Çalışma Sayfası (Sheet)", sheet_names)
                    df_loaded = pd.read_excel(uploaded_file, sheet_name=selected_sheet)
                
                st.session_state.raw_df = df_loaded.copy()
                st.session_state.df = df_loaded.copy()
                log_action(f"'{uploaded_file.name}' dosyası başarıyla yüklendi ({len(df_loaded)} satır, {len(df_loaded.columns)} sütun).")
                st.success("Dosya başarıyla yüklendi!")
            except Exception as e:
                st.error(f"Dosya okunurken bir hata oluştu: {str(e)}")

        if st.session_state.df is not None:
            st.divider()
            st.subheader("🔄 Hızlı İşlemler")
            if st.button("⏪ Orijinal Veriye Sıfırla", use_container_width=True):
                st.session_state.df = st.session_state.raw_df.copy()
                log_action("Veri seti ilk yüklenen haline sıfırlandı.")
                st.rerun()

            st.caption(f"Aktif Dosya: **{uploaded_file.name}**")

# -------------------------------------------------------------
# ANA EKRAN BAŞLIĞI
# -------------------------------------------------------------
st.markdown('<div class="main-header">📊 Excel & CSV Veri Temizleme ve Özetleme Aracı</div>', unsafe_allow_html=True)
st.markdown('<div class="sub-header">Verilerinizi yükleyin, hızlıca analiz edin, temizleyin ve temiz halini Excel/CSV olarak indirin.</div>', unsafe_allow_html=True)

if st.session_state.df is None:
    st.info("👈 Başlamak için lütfen sol taraftaki panelden bir **Excel (.xlsx, .xls)** veya **CSV** dosyası yükleyin.")
    st.stop()

# Aktif DataFrame referansı
df = st.session_state.df

# -------------------------------------------------------------
# ANA SEKME YAPISI
# -------------------------------------------------------------
tab_summary, tab_cleaning, tab_export = st.tabs([
    "📈 1. Veri Özeti & Analiz", 
    "🧹 2. Veri Temizleme", 
    "💾 3. Dışa Aktar & İndir"
])

# =============================================================
# SEKME 1: VERİ ÖZETİ VE PROFİLLEME
# =============================================================
with tab_summary:
    st.subheader("📌 Genel Bakış & Temel Metrikler")
    
    total_rows = len(df)
    total_cols = len(df.columns)
    total_cells = total_rows * total_cols
    total_missing = df.isnull().sum().sum()
    missing_pct = (total_missing / total_cells * 100) if total_cells > 0 else 0
    duplicate_rows = df.duplicated().sum()
    memory_usage_mb = df.memory_usage(deep=True).sum() / (1024 * 1024)

    col1, col2, col3, col4, col5 = st.columns(5)
    col1.metric("Toplam Satır", f"{total_rows:,}")
    col2.metric("Toplam Sütun", total_cols)
    col3.metric("Eksik Değerler", f"{total_missing:,}", f"{missing_pct:.1f}%")
    col4.metric("Yinelenen Satırlar", f"{duplicate_rows:,}")
    col5.metric("Bellek Kullanımı", f"{memory_usage_mb:.2f} MB")

    st.divider()

    # Veri Önizlemesi
    st.subheader("👀 Veri Önizlemesi")
    view_option = st.radio("Görünüm Modu:", ["İlk 10 Satır", "Son 10 Satır", "Rastgele 10 Satır", "Tüm Veri"], horizontal=True)
    if view_option == "İlk 10 Satır":
        st.dataframe(df.head(10), use_container_width=True)
    elif view_option == "Son 10 Satır":
        st.dataframe(df.tail(10), use_container_width=True)
    elif view_option == "Rastgele 10 Satır":
        st.dataframe(df.sample(min(10, total_rows)), use_container_width=True)
    else:
        st.dataframe(df, use_container_width=True)

    st.divider()

    # Sütun İnceleme ve Veri Tipleri
    st.subheader("📋 Sütun Sağlığı ve Veri Tipleri")
    col_info = []
    for col in df.columns:
        null_count = df[col].isnull().sum()
        col_info.append({
            "Sütun Adı": col,
            "Veri Tipi": str(df[col].dtype),
            "Dolu Satır": total_rows - null_count,
            "Eksik Satır": null_count,
            "Eksiklik Oranı (%)": round((null_count / total_rows) * 100, 2) if total_rows > 0 else 0,
            "Benzersiz Değer": df[col].nunique(),
            "Örnek Değer": str(df[col].dropna().iloc[0]) if not df[col].dropna().empty else "-"
        })
    info_df = pd.DataFrame(col_info)
    st.dataframe(info_df, use_container_width=True)

    # Eksik Değer Dağılım Grafiği
    missing_cols = info_df[info_df["Eksik Satır"] > 0]
    if not missing_cols.empty:
        st.subheader("📉 Sütunlardaki Eksik Değer Dağılımı")
        chart_data = missing_cols.set_index("Sütun Adı")["Eksiklik Oranı (%)"]
        st.bar_chart(chart_data)
    else:
        st.success("🎉 Tebrikler! Veri setinde hiç eksik değer (null/NaN) bulunmuyor.")

    # İstatistiksel Özet
    st.divider()
    st.subheader("📊 İstatistiksel Dağılım (Describe)")
    describe_type = st.radio("İstatistik Kapsamı:", ["Sayısal Sütunlar", "Tüm Sütunlar"], horizontal=True)
    if describe_type == "Sayısal Sütunlar":
        num_df = df.select_dtypes(include='number')
        if not num_df.empty:
            st.dataframe(num_df.describe().T, use_container_width=True)
        else:
            st.info("Veri setinde sayısal sütun bulunamadı.")
    else:
        st.dataframe(df.describe(include='all').astype(str).T, use_container_width=True)


# =============================================================
# SEKME 2: VERİ TEMİZLEME ARAÇLARI
# =============================================================
with tab_cleaning:
    st.subheader("🛠️ Temizleme & Veri Düzenleme Araçları")

    # 1. YİNELENEN VERİLER (DUPLICATES)
    with st.expander("👯 1. Yinelenen (Duplicate) Satırları Temizleme", expanded=False):
        dup_count = df.duplicated().sum()
        st.write(f"Mevcut Tam Yinelenen Satır Sayısı: **{dup_count}**")

        dup_mode = st.radio(
            "Yinelenme Kriteri:",
            ["Tüm Sütunlar Aynı Olanlar", "Belirli Sütunları Aynı Olanlar"],
            key="dup_mode"
        )
        subset_cols = None
        if dup_mode == "Belirli Sütunları Aynı Olanlar":
            subset_cols = st.multiselect("Kriter Sütunları Seçin:", df.columns.tolist())

        keep_option = st.selectbox("Hangi Kayıt Saklansın?", ["İlk Kayıt (first)", "Son Kayıt (last)"], index=0)
        keep_val = "first" if "İlk" in keep_option else "last"

        if st.button("🗑️ Yinelenen Satırları Temizle", type="primary"):
            cols_to_check = subset_cols if (subset_cols and len(subset_cols) > 0) else None
            before_len = len(df)
            st.session_state.df = df.drop_duplicates(subset=cols_to_check, keep=keep_val).reset_index(drop=True)
            removed = before_len - len(st.session_state.df)
            log_action(f"{removed} adet yinelenen satır temizlendi (Saklanan: {keep_val}).")
            st.success(f"{removed} adet satır başarıyla silindi!")
            st.rerun()

    # 2. EKSİK DEĞERLER (MISSING VALUES)
    with st.expander("❓ 2. Eksik Değer (NaN / Null) Yönetimi", expanded=False):
        st.markdown("**A) Eksik Değer İçeren Satırları Silme**")
        drop_na_mode = st.selectbox(
            "Silme Koşulu:",
            ["Tüm satır boşsa sil (all)", "Herhangi bir hücre boşsa sil (any)", "Seçili sütunlar boşsa sil (subset)"]
        )
        drop_subset = None
        if "Seçili" in drop_na_mode:
            drop_subset = st.multiselect("Kontrol edilecek sütunlar:", df.columns.tolist(), key="dropna_cols")

        if st.button("❌ Eksik Satırları Sil"):
            before_len = len(df)
            how_val = "all" if "all" in drop_na_mode else "any"
            st.session_state.df = df.dropna(how=how_val, subset=drop_subset if drop_subset else None).reset_index(drop=True)
            removed = before_len - len(st.session_state.df)
            log_action(f"Eksik veri içeren {removed} satır silindi (Mod: {drop_na_mode}).")
            st.success(f"{removed} adet satır silindi.")
            st.rerun()

        st.divider()
        st.markdown("**B) Eksik Değerleri Doldurma (Imputation)**")
        fill_col = st.selectbox("Doldurulacak Sütun:", df.columns.tolist(), key="fill_col")
        
        col_type = df[fill_col].dtype
        null_in_col = df[fill_col].isnull().sum()
        st.caption(f"'{fill_col}' sütununda **{null_in_col}** adet eksik değer var. (Tip: {col_type})")

        fill_methods = ["Sabit Değer", "En Çok Tekrar Eden Değer (Mod)", "Önceki Değerle Doldur (ffill)", "Sonraki Değerle Doldur (bfill)"]
        if pd.api.types.is_numeric_dtype(df[fill_col]):
            fill_methods = ["Ortalama (Mean)", "Medyan (Median)"] + fill_methods

        chosen_method = st.selectbox("Doldurma Yöntemi:", fill_methods)
        custom_fill_value = None
        if chosen_method == "Sabit Değer":
            custom_fill_value = st.text_input("Doldurulacak Değer:", value="0" if pd.api.types.is_numeric_dtype(df[fill_col]) else "Bilinmiyor")

        if st.button("✏️ Eksik Değerleri Doldur"):
            if null_in_col == 0:
                st.info("Bu sütunda zaten eksik değer yok.")
            else:
                updated_col = df[fill_col].copy()
                if chosen_method == "Ortalama (Mean)":
                    val = updated_col.mean()
                    updated_col.fillna(val, inplace=True)
                elif chosen_method == "Medyan (Median)":
                    val = updated_col.median()
                    updated_col.fillna(val, inplace=True)
                elif chosen_method == "En Çok Tekrar Eden Değer (Mod)":
                    mode_val = updated_col.mode()
                    val = mode_val[0] if not mode_val.empty else ""
                    updated_col.fillna(val, inplace=True)
                elif chosen_method == "Önceki Değerle Doldur (ffill)":
                    updated_col.ffill(inplace=True)
                    val = "ffill"
                elif chosen_method == "Sonraki Değerle Doldur (bfill)":
                    updated_col.bfill(inplace=True)
                    val = "bfill"
                elif chosen_method == "Sabit Değer":
                    if pd.api.types.is_numeric_dtype(df[fill_col]):
                        try:
                            val = float(custom_fill_value) if "." in custom_fill_value else int(custom_fill_value)
                        except ValueError:
                            val = custom_fill_value
                    else:
                        val = custom_fill_value
                    updated_col.fillna(val, inplace=True)

                st.session_state.df[fill_col] = updated_col
                log_action(f"'{fill_col}' sütunundaki eksik değerler '{chosen_method}' yöntemiyle dolduruldu.")
                st.success(f"'{fill_col}' sütunu başarıyla dolduruldu.")
                st.rerun()

    # 3. SÜTUN YÖNETİMİ (COLUMNS)
    with st.expander("📑 3. Sütun Yönetimi (Silme, Yeniden Adlandırma & Normalizasyon)", expanded=False):
        c1, c2 = st.columns(2)
        with c1:
            st.markdown("**Gereksiz Sütunları Kaldırma**")
            cols_to_drop = st.multiselect("Silinecek Sütunları Seçin:", df.columns.tolist())
            if st.button("🗑️ Seçili Sütunları Sil", disabled=len(cols_to_drop) == 0):
                st.session_state.df.drop(columns=cols_to_drop, inplace=True)
                log_action(f"Sütunlar silindi: {', '.join(cols_to_drop)}")
                st.success(f"{len(cols_to_drop)} sütun silindi.")
                st.rerun()

        with c2:
            st.markdown("**Sütun İsimlerini Normalize Etme**")
            st.caption("Boşlukları alt çizgiye dönüştürür, küçük harf yapar, Türkçe karakterleri standartlaştırır.")
            if st.button("✨ Sütun İsimlerini Otomatik Düzelt"):
                tr_map = str.maketrans("çğıöşüÇĞİÖŞÜ ", "cgiosuCGIOSU_")
                new_cols = [c.strip().translate(tr_map).lower() for c in df.columns]
                st.session_state.df.columns = new_cols
                log_action("Sütun adları normalize edildi (snake_case / tr karakterler temizlendi).")
                st.success("Sütun adları başarıyla düzenlendi.")
                st.rerun()

        st.divider()
        st.markdown("**Tek Bir Sütunu Yeniden Adlandırma**")
        col_to_rename = st.selectbox("Adı değiştirilecek sütun:", df.columns.tolist(), key="rename_col")
        new_name = st.text_input("Yeni İsim:", value=col_to_rename)
        if st.button("Adı Güncelle", disabled=(new_name == col_to_rename or not new_name.strip())):
            st.session_state.df.rename(columns={col_to_rename: new_name.strip()}, inplace=True)
            log_action(f"'{col_to_rename}' sütunu '{new_name.strip()}' olarak yeniden adlandırıldı.")
            st.success(f"Sütun adı güncellendi: {new_name.strip()}")
            st.rerun()

    # 4. VERİ TİPİ DÖNÜŞTÜRME
    with st.expander("🔄 4. Veri Tipi Dönüştürme (Type Casting)", expanded=False):
        c_col, c_type = st.columns(2)
        with c_col:
            target_col = st.selectbox("Dönüştürülecek Sütun:", df.columns.tolist(), key="cast_col")
        with c_type:
            target_type = st.selectbox("Hedef Veri Tipi:", ["Sayısal (Numeric / Float)", "Tam Sayı (Integer)", "Metin (String / Text)", "Tarih/Saat (Datetime)", "Kategori (Category)"])

        if st.button("Dönüşümü Uygula"):
            try:
                if target_type == "Sayısal (Numeric / Float)":
                    st.session_state.df[target_col] = pd.to_numeric(st.session_state.df[target_col], errors='coerce')
                elif target_type == "Tam Sayı (Integer)":
                    st.session_state.df[target_col] = pd.to_numeric(st.session_state.df[target_col], errors='coerce').fillna(0).astype('int64')
                elif target_type == "Metin (String / Text)":
                    st.session_state.df[target_col] = st.session_state.df[target_col].astype(str)
                elif target_type == "Tarih/Saat (Datetime)":
                    st.session_state.df[target_col] = pd.to_datetime(st.session_state.df[target_col], errors='coerce')
                elif target_type == "Kategori (Category)":
                    st.session_state.df[target_col] = st.session_state.df[target_col].astype('category')

                log_action(f"'{target_col}' sütunu '{target_type}' tipine dönüştürüldü.")
                st.success(f"'{target_col}' tipi başarıyla güncellendi!")
                st.rerun()
            except Exception as e:
                st.error(f"Dönüştürme hatası: {str(e)}")

    # 5. METİN DÜZENLEME (STRING CLEANING)
    with st.expander("🔤 5. Metin Temizleme (Boşluk Temizleme & Büyük/Küçük Harf)", expanded=False):
        text_cols = df.select_dtypes(include=['object', 'string']).columns.tolist()
        if not text_cols:
            st.info("Veri setinde metin (string/object) sütunu bulunmuyor.")
        else:
            selected_text_col = st.selectbox("Metin Sütunu:", text_cols, key="text_col_select")
            text_action = st.radio(
                "Uygulanacak Metin İşlemi:",
                ["Baş/Son Boşlukları Kırp (Strip Whitespace)", "Küçük Harf Yap (lowercase)", "BÜYÜK HARF YAP (UPPERCASE)", "Baş Harfleri Büyüt (Title Case)"],
                horizontal=True
            )
            if st.button("Metin İşlemini Uygula"):
                series = df[selected_text_col].astype(str)
                if text_action == "Baş/Son Boşlukları Kırp (Strip Whitespace)":
                    st.session_state.df[selected_text_col] = series.str.strip()
                elif text_action == "Küçük Harf Yap (lowercase)":
                    st.session_state.df[selected_text_col] = series.str.lower()
                elif text_action == "BÜYÜK HARF YAP (UPPERCASE)":
                    st.session_state.df[selected_text_col] = series.str.upper()
                elif text_action == "Baş Harfleri Büyüt (Title Case)":
                    st.session_state.df[selected_text_col] = series.str.title()

                log_action(f"'{selected_text_col}' sütununa '{text_action}' uygulandı.")
                st.success("Metin düzenleme tamamlandı.")
                st.rerun()

    # 6. AYKIRI DEĞER (OUTLIER) FİLTRESİ
    with st.expander("📐 6. Aykırı Değer Filtreleme (Sayısal Sütunlar)", expanded=False):
        num_cols = df.select_dtypes(include='number').columns.tolist()
        if not num_cols:
            st.info("Aykırı değer tespiti için sayısal sütun bulunmuyor.")
        else:
            selected_num_col = st.selectbox("Sayısal Sütun Seçin:", num_cols, key="outlier_col")
            outlier_mode = st.radio("Yöntem:", ["IQR (Çeyrekler Açıklığı - Önerilen)", "Özel Min / Max Eşiği"], horizontal=True)

            if outlier_mode == "IQR (Çeyrekler Açıklığı - Önerilen)":
                q1 = df[selected_num_col].quantile(0.25)
                q3 = df[selected_num_col].quantile(0.75)
                iqr = q3 - q1
                factor = st.slider("IQR Çarpanı (Varsayılan: 1.5):", min_value=1.0, max_value=3.0, value=1.5, step=0.1)
                lower_bound = q1 - factor * iqr
                upper_bound = q3 + factor * iqr
                outliers_count = ((df[selected_num_col] < lower_bound) | (df[selected_num_col] > upper_bound)).sum()
                st.write(f"Alt Sınır: **{lower_bound:.2f}**, Üst Sınır: **{upper_bound:.2f}** | Tespit Edilen Aykırı Satır: **{outliers_count}**")

                if st.button("Aykırı Değerleri Filtrele (Satırları Sil)", disabled=outliers_count == 0):
                    before_len = len(df)
                    st.session_state.df = df[(df[selected_num_col] >= lower_bound) & (df[selected_num_col] <= upper_bound)].reset_index(drop=True)
                    removed = before_len - len(st.session_state.df)
                    log_action(f"'{selected_num_col}' sütunundaki {removed} adet aykırı değer silindi (IQR={factor}).")
                    st.success(f"{removed} adet aykırı satır filtrelendi.")
                    st.rerun()
            else:
                col_min = float(df[selected_num_col].min()) if not df[selected_num_col].empty else 0.0
                col_max = float(df[selected_num_col].max()) if not df[selected_num_col].empty else 100.0
                c_min, c_max = st.columns(2)
                with c_min:
                    min_val = st.number_input("Minimum Değer:", value=col_min)
                with c_max:
                    max_val = st.number_input("Maksimum Değer:", value=col_max)

                if st.button("Eşik Değerlere Göre Filtrele"):
                    before_len = len(df)
                    st.session_state.df = df[(df[selected_num_col] >= min_val) & (df[selected_num_col] <= max_val)].reset_index(drop=True)
                    removed = before_len - len(st.session_state.df)
                    log_action(f"'{selected_num_col}' sütununa [{min_val}, {max_val}] eşik filtresi uygulandı ({removed} satır elendi).")
                    st.success(f"{removed} satır filtrelendi.")
                    st.rerun()


# =============================================================
# SEKME 3: DIŞA AKTARMA VE İNDİRME
# =============================================================
with tab_export:
    st.subheader("💾 Temizlenmiş Veriyi İndirme")
    
    st.write(f"Mevcut Veri Boyutu: **{len(df):,} Satır** × **{len(df.columns)} Sütun**")
    st.dataframe(df.head(10), use_container_width=True)

    col_dl1, col_dl2 = st.columns(2)

    # EXCEL İNDİRME (.xlsx)
    with col_dl1:
        st.markdown("### 📗 Excel (.xlsx)")
        st.caption("openpyxl ile stillendirilmiş, otomatik sütun genişlikli dosya.")
        with st.spinner("Excel dosyası hazırlanıyor..."):
            excel_data = to_excel_bytes(df)
        
        default_excel_name = f"temiz_veri_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        st.download_button(
            label="📥 Excel (.xlsx) Olarak İndir",
            data=excel_data,
            file_name=default_excel_name,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            use_container_width=True
        )

    # CSV İNDİRME (.csv)
    with col_dl2:
        st.markdown("### 📄 CSV (.csv)")
        st.caption("Excel ile uyumlu (UTF-8 BOM) ve ayraç seçilebilir CSV.")
        csv_sep = st.selectbox("CSV Ayırıcı Karakter:", [",", ";", "\\t"], index=0)
        
        # utf-8-sig Türkçe karakterlerin Excel'de düzgün açılmasını sağlar
        csv_data = df.to_csv(index=False, sep=csv_sep).encode('utf-8-sig')
        default_csv_name = f"temiz_veri_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        
        st.download_button(
            label="📥 CSV (.csv) Olarak İndir",
            data=csv_data,
            file_name=default_csv_name,
            mime="text/csv",
            use_container_width=True
        )

    st.divider()

    # İŞLEM GEÇMİŞİ (AUDIT / CHANGE LOG)
    st.subheader("📜 Yapılan İşlemlerin Geçmişi (Audit Log)")
    if st.session_state.history_log:
        for log_entry in reversed(st.session_state.history_log):
            st.markdown(f"- {log_entry}")
    else:
        st.info("Henüz herhangi bir temizleme işlemi uygulanmadı.")
